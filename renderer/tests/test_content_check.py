from pptx import Presentation
from pptx.util import Inches

from renderer.content_check import missing_content


def test_missing_table_cannot_pass_shape_count_quality(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(9), Inches(1)).text = "Quarterly results"
    output = tmp_path / "missing.pptx"
    prs.save(output)
    html = "<html><body><h1>Quarterly results</h1><table><tr><td>Q2</td><td>$3.9M</td></tr></table></body></html>"
    assert missing_content([html], output)[0]["missing_words"] == 3
    table = slide.shapes.add_table(1, 2, 0, Inches(2), Inches(9), Inches(1)).table
    table.cell(0, 0).text = "Q2"
    table.cell(0, 1).text = "$3.9M"
    prs.save(output)
    assert missing_content([html], output) == []


def test_formatting_inside_a_word_and_head_metadata_are_not_content_loss(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(9), Inches(1)).text = "Example Retail\nGroup"
    output = tmp_path / "styled.pptx"
    prs.save(output)
    assert (
        missing_content(
            [
                "<html><head><title>Not slide content</title><style>p{color:red}</style></head><body><p>Ex<strong>ample</strong> Retail</p><p>Group</p></body></html>"
            ],
            output,
        )
        == []
    )
