"""Check authored words against the delivered PPTX, before any upload.

Editability counts start after classification. This independent comparison
also catches a whole table disappearing before an emit operation was created.
"""

from collections import Counter
from html.parser import HTMLParser
from pathlib import Path
import re
import unicodedata

from pptx import Presentation


class _Text(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"head", "script", "style"}:
            self.skip += 1
        elif tag not in {"span", "em", "strong", "b", "i", "a", "small", "sup", "sub"}:
            self.parts.append(" ")

    def handle_endtag(self, tag):
        if tag in {"head", "script", "style"}:
            self.skip = max(0, self.skip - 1)
        elif tag not in {"span", "em", "strong", "b", "i", "a", "small", "sup", "sub"}:
            self.parts.append(" ")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def _words(text: str) -> Counter:
    return Counter(re.findall(r"\w+", unicodedata.normalize("NFKC", text).casefold()))


def _shape_text(shapes):
    for shape in shapes:
        if shape.has_text_frame:
            yield shape.text
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    if not cell.is_spanned:
                        yield cell.text
        if hasattr(shape, "shapes"):
            yield from _shape_text(shape.shapes)


def missing_content(slides: list[str], pptx: Path) -> list[dict]:
    prs = Presentation(pptx)
    failures = []
    if len(prs.slides) != len(slides):
        return [{"error": f"expected {len(slides)} slides, exported {len(prs.slides)}"}]
    for index, (html, slide) in enumerate(zip(slides, prs.slides, strict=True)):
        source = _Text()
        source.feed(html)
        expected = _words("".join(source.parts))
        actual = _words(" ".join(_shape_text(slide.shapes)))
        missing = expected - actual
        if missing:
            failures.append(
                {
                    "slide_index": index,
                    "missing_words": sum(missing.values()),
                    "sample": list(missing)[:12],
                }
            )
    return failures
