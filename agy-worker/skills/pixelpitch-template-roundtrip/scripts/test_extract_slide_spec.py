"""Template styles belong to the OOXML inheritance chain, not just runs."""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from extract_slide_spec import extract


def test_resolves_layout_run_style_and_normalizes_to_canvas(tmp_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(26.6666667)
    prs.slide_height = Inches(15)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "An inherited title"
    layout_title = title._base_placeholder
    styles = layout_title._element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}txBody"
    ).find("{http://schemas.openxmlformats.org/drawingml/2006/main}lstStyle")
    level = OxmlElement("a:lvl1pPr")
    defaults = OxmlElement("a:defRPr")
    defaults.set("sz", "6000")
    defaults.set("b", "1")
    latin = OxmlElement("a:latin")
    latin.set("typeface", "Roboto Medium")
    defaults.append(latin)
    fill = OxmlElement("a:solidFill")
    color = OxmlElement("a:srgbClr")
    color.set("val", "1971ED")
    fill.append(color)
    defaults.append(fill)
    level.append(defaults)
    styles.append(level)
    source = tmp_path / "template.pptx"
    prs.save(source)

    extract(source, tmp_path / "spec-out", base_renders=False)

    import json

    record = json.loads((tmp_path / "spec-out/spec/slide-01.json").read_text())
    run = record["shapes"][0]["text"][0]["runs"][0]
    assert run["font"] == "Roboto Medium"
    assert run["size_pt"] == 60
    assert run["size_px"] == pytest.approx(40, abs=0.01)
    assert run["color"] == "#1971ED"
    assert run["bold"] is True


def test_direct_style_overrides_layout_and_table_text_is_extracted(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Direct title"
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(31)
    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(8), Inches(2)
    ).table
    table.cell(0, 0).text = "Quarter"
    table.cell(1, 1).text = "$3.9M"
    source = tmp_path / "template.pptx"
    prs.save(source)
    extract(source, tmp_path / "out", base_renders=False)
    import json

    record = json.loads((tmp_path / "out/spec/slide-01.json").read_text())
    assert record["shapes"][0]["text"][0]["runs"][0]["size_pt"] == 31
    extracted = next(shape for shape in record["shapes"] if "table" in shape)
    assert extracted["table"]["rows"][1][1]["text"][0]["runs"][0]["text"] == "$3.9M"
