#!/usr/bin/env python3
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     https://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Extract a slide-by-slide construction spec from a source PPTX.

Mechanical extraction only — the *reconstruction* is the harness's job. For
every slide this writes ``spec/slide-NN.json`` with the full shape tree
(geometry in 1280x720 px space, fills, strokes, text runs with exact font
styles, picture blobs saved under ``spec/media/``), plus ``base/slide-NN.png``
— a pixel-exact 1280-wide render of the slide for image-base reconstruction
or visual comparison.

Run with the renderer venv python (python-pptx):
  renderer/.venv/bin/python extract_slide_spec.py --template template.pptx --outdir spec-out
"""

from __future__ import annotations

import argparse
import json
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

TARGET_W = 1280


def _fill_hex(shape) -> dict:
    try:
        fill = shape.fill
    except (AttributeError, ValueError):
        return {"fill": None}
    try:
        if fill.type is None:
            return {"fill": "none"}
        if str(fill.type).startswith("SOLID"):
            color = fill.fore_color
            if color.type is not None and str(color.type).startswith("SCHEME"):
                return {"fill": {"scheme": str(color.theme_color)}}
            return {"fill": f"#{color.rgb}"}
        if str(fill.type).startswith("GRADIENT"):
            stops = []
            try:
                for stop in fill.gradient_stops:
                    stops.append(f"#{stop.color.rgb}@{stop.position:.2f}")
            except Exception:
                pass
            return {"fill": {"gradient": stops}}
        if str(fill.type).startswith("BACKGROUND"):
            return {"fill": "none"}
    except Exception:
        pass
    return {"fill": None}


def _line(shape) -> dict:
    try:
        line = shape.line
        if line.fill.type is None or str(line.fill.type).startswith("BACKGROUND"):
            return {"line": None}
        color = line.color
        if color and color.type is not None and str(color.type).startswith("SCHEME"):
            return {"line": {"scheme": str(color.theme_color)}}
        weight = line.width.pt if line.width else None
        return {"line": {"hex": f"#{color.rgb}", "pt": weight}}
    except Exception:
        return {"line": None}


def _theme(slide) -> dict:
    master = slide.slide_layout.slide_master
    root = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    colors = {}
    for node in root.findall(".//a:clrScheme/*", ns):
        color = node[0]
        colors[etree.QName(node).localname] = color.get("lastClr") or color.get("val")
    fonts = {}
    for kind, prefix in (("majorFont", "+mj"), ("minorFont", "+mn")):
        for slot, suffix in (("latin", "lt"), ("ea", "ea"), ("cs", "cs")):
            node = root.find(f".//a:{kind}/a:{slot}", ns)
            if node is not None:
                fonts[f"{prefix}-{suffix}"] = node.get("typeface")
    mapping = master._element.find(qn("p:clrMap"))
    color_map = dict(mapping.attrib) if mapping is not None else {}
    for owner in (slide.slide_layout, slide):
        override = owner._element.find("./" + qn("p:clrMapOvr") + "/" + qn("a:overrideClrMapping"))
        if override is not None:
            color_map.update(override.attrib)
    return {"colors": colors, "fonts": fonts, "color_map": color_map}


def _color(node, theme: dict) -> str | None:
    if node is None or len(node) == 0:
        return None
    color = node[0]
    kind = etree.QName(color).localname
    value = color.get("val")
    if kind == "schemeClr":
        value = theme["colors"].get(theme["color_map"].get(value, value))
    elif kind == "sysClr":
        value = color.get("lastClr")
    if not value or len(value) != 6:
        return None
    channels = [int(value[i:i + 2], 16) for i in (0, 2, 4)]
    for transform in color:
        amount = int(transform.get("val", "100000")) / 100000
        name = etree.QName(transform).localname
        if name in {"shade", "lumMod"}:
            channels = [c * amount for c in channels]
        elif name in {"tint", "lumOff"}:
            channels = [c + (255 - c if name == "tint" else 255) * amount for c in channels]
    return "#" + "".join(f"{max(0, min(255, round(c))):02X}" for c in channels)


def _paragraph_sources(shape, para, slide) -> list:
    """Least to most specific, including placeholders' list-level defaults."""
    level = f"a:lvl{para.level + 1}pPr"
    kind = "otherStyle"
    if shape is not None and shape.is_placeholder:
        role = str(shape.placeholder_format.type)
        kind = "titleStyle" if "TITLE" in role and "SUBTITLE" not in role else "bodyStyle" if "BODY" in role else "otherStyle"
    sources = []
    master_style = slide.slide_layout.slide_master._element.find(
        "./" + qn("p:txStyles") + "/" + qn("p:" + kind)
    )
    if master_style is not None:
        sources.extend([master_style.find(qn("a:defPPr")), master_style.find(qn(level))])
    chain = []
    current = shape
    while current is not None:
        chain.append(current)
        current = getattr(current, "_base_placeholder", None)
    for owner in reversed(chain):
        body = owner._element.find(qn("p:txBody"))
        if body is not None:
            styles = body.find(qn("a:lstStyle"))
            if styles is not None:
                sources.extend([styles.find(qn("a:defPPr")), styles.find(qn(level))])
    sources.append(para._p.find(qn("a:pPr")))
    return [node for node in sources if node is not None]


def _runs(frame, shape, slide, scale: float, theme: dict) -> list[dict]:
    paragraphs = []
    for para in frame.paragraphs:
        sources = _paragraph_sources(shape, para, slide)
        p_attrs = {}
        for node in sources:
            p_attrs.update(node.attrib)
        runs = []
        for run_element in para._p:
            if run_element.tag not in {qn("a:r"), qn("a:fld")}:
                continue
            text_element = run_element.find(qn("a:t"))
            text_value = text_element.text or "" if text_element is not None else ""
            attrs, children = {}, {}
            defaults = [node.find(qn("a:defRPr")) for node in sources]
            for node in [*defaults, run_element.find(qn("a:rPr"))]:
                if node is not None:
                    attrs.update(node.attrib)
                    children.update({child.tag: child for child in node})
            latin = children.get(qn("a:latin"))
            font = latin.get("typeface") if latin is not None else "+mn-lt"
            font = theme["fonts"].get(font, font)
            size = int(attrs["sz"]) / 100 if "sz" in attrs else None
            runs.append(
                {
                    "text": text_value,
                    "font": font,
                    "size_pt": size,
                    "size_px": round(size * 12700 * scale, 3) if size else None,
                    "bold": attrs.get("b", "0") in {"1", "true"},
                    "italic": attrs.get("i", "0") in {"1", "true"},
                    "color": _color(children.get(qn("a:solidFill")), theme),
                }
            )
        spacing = {}
        for node in sources:
            for key in ("lnSpc", "spcBef", "spcAft"):
                value = node.find(qn("a:" + key))
                if value is not None and len(value):
                    item = value[0]
                    spacing[key] = ({"multiple": int(item.get("val")) / 100000}
                                    if item.tag == qn("a:spcPct") else
                                    {"px": int(item.get("val")) / 100 * 12700 * scale})
        paragraphs.append(
            {
                "align": {"l": "LEFT", "r": "RIGHT", "ctr": "CENTER", "just": "JUSTIFY"}.get(p_attrs.get("algn"), "LEFT"),
                "level": para.level,
                "runs": runs,
                "spacing": spacing,
            }
        )
    return paragraphs


def _shape_entry(shape, scale: float, media_dir: Path, slide_no: int, counter: list, slide, theme: dict) -> dict:
    entry: dict = {
        "kind": str(shape.shape_type).split()[0],
        "name": shape.name,
        "shape_id": shape.shape_id,
        "x": round(shape.left * scale, 1) if shape.left is not None else None,
        "y": round(shape.top * scale, 1) if shape.top is not None else None,
        "w": round(shape.width * scale, 1) if shape.width is not None else None,
        "h": round(shape.height * scale, 1) if shape.height is not None else None,
        "rotation": shape.rotation or 0,
    }
    if shape._element.xpath('.//a:fld[@type="slidenum"]'):
        entry["slide_number_field"] = True
    if shape.is_placeholder:
        try:
            entry["placeholder"] = str(shape.placeholder_format.type).split()[0]
        except Exception:
            pass
    entry.update(_fill_hex(shape))
    entry.update(_line(shape))
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        entry["text"] = _runs(shape.text_frame, shape, slide, scale, theme)
        entry["insets"] = [round(getattr(shape.text_frame, "margin_" + side) * scale, 3)
                           for side in ("top", "right", "bottom", "left")]
        current = shape
        while current is not None:
            props = current._element.find("./" + qn("p:txBody") + "/" + qn("a:bodyPr"))
            if props is not None and props.get("anchor"):
                entry["anchor"] = props.get("anchor")
                break
            current = getattr(current, "_base_placeholder", None)
    if getattr(shape, "has_table", False):
        table = shape.table
        entry["table"] = {
            "widths": [round(column.width * scale, 3) for column in table.columns],
            "heights": [round(row.height * scale, 3) for row in table.rows],
            "rows": [[{"text": _runs(cell.text_frame, None, slide, scale, theme),
                       "insets": [round(getattr(cell, "margin_" + side) * scale, 3) for side in ("top", "right", "bottom", "left")],
                       "rowspan": cell.span_height, "colspan": cell.span_width,
                       "spanned": cell.is_spanned}
                      for cell in row.cells] for row in table.rows],
        }
    if entry["kind"] == "PICTURE":
        try:
            image = shape.image
            counter[0] += 1
            ext = image.ext or "png"
            fname = f"s{slide_no:02d}-img{counter[0]:02d}.{ext}"
            (media_dir / fname).write_bytes(image.blob)
            entry["image"] = str((media_dir / fname).name)
        except Exception:
            entry["image"] = None
    if entry["kind"] == "GROUP":
        entry["children"] = [
            _shape_entry(child, scale, media_dir, slide_no, counter, slide, theme)
            for child in shape.shapes
        ]
    return entry


def extract(template: Path, outdir: Path, base_renders: bool = True) -> dict:
    prs = Presentation(str(template))
    slide_width = int(prs.slide_width or 12192000)  # 12192000 EMU = 13.33in
    scale = TARGET_W / slide_width
    spec_dir = outdir / "spec"
    media_dir = spec_dir / "media"
    base_dir = outdir / "base"
    media_dir.mkdir(parents=True, exist_ok=True)
    base_dir.mkdir(parents=True, exist_ok=True)

    index = []
    for i, slide in enumerate(prs.slides, start=1):
        counter = [0]
        theme = _theme(slide)
        entries = [
            _shape_entry(shape, scale, media_dir, i, counter, slide, theme)
            for shape in slide.shapes
        ]
        record = {
            "schema_version": 2,
            "slide": i,
            "layout": slide.slide_layout.name,
            "canvas": [TARGET_W, round(prs.slide_height * scale)],
            "source_size_emu": [int(prs.slide_width), int(prs.slide_height)],
            "font_scale_px_per_pt": 12700 * scale,
            "shapes": entries,
        }
        # Layout/master page-number fields must not become the source page's
        # literal number in every newly authored slide's background image.
        number_shape = next((shape for owner in (slide, slide.slide_layout, slide.slide_layout.slide_master)
                             for shape in owner.shapes if shape._element.xpath('.//a:fld[@type="slidenum"]')), None)
        if number_shape is not None:
            record["slide_number"] = _shape_entry(number_shape, scale, media_dir, i, counter, slide, theme)
        (spec_dir / f"slide-{i:02d}.json").write_text(
            json.dumps(record, indent=1), encoding="utf-8"
        )
        index.append(
            {
                "slide": i,
                "layout": record["layout"],
                "shapes": len(entries),
                "pictures": sum(1 for e in entries if e["kind"] == "PICTURE"),
                "text_shapes": sum(1 for e in entries if "text" in e),
            }
        )

    # Pixel-exact base renders for image-base mode and visual comparison.
    if not base_renders:
        return {
            "slides": index,
            "canvas": [TARGET_W, round(prs.slide_height * scale)],
            "base_renders": "skipped",
        }
    with tempfile.TemporaryDirectory(prefix="template-lo-") as profile:
        subprocess.run(
            ["soffice", f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(template)],
            check=True, capture_output=True, timeout=420,
        )
    pdf = outdir / (template.stem + ".pdf")
    if not pdf.is_file():
        raise RuntimeError("template extraction produced no PDF")
    if pdf.is_file():
        subprocess.run(
            ["pdftoppm", "-png", "-scale-to", str(TARGET_W), str(pdf), str(base_dir / "slide")],
            check=True, capture_output=True, timeout=420,
        )
        pdf.unlink()
    for number in range(1, len(prs.slides) + 1):
        # pdftoppm pads to the page count; the bundle uses two-digit names.
        rendered = next((p for p in base_dir.glob("slide-*.png") if int(p.stem.split("-")[-1]) == number), None)
        if rendered is None:
            raise RuntimeError(f"template extraction is missing rendered slide {number}")
        rendered.rename(base_dir / f"slide-{number:02d}.png")
    return {"slides": index, "canvas": [TARGET_W, round(prs.slide_height * scale)]}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", required=True)
    parser.add_argument("--outdir", required=True)
    parser.add_argument(
        "--no-base",
        action="store_true",
        help="extract geometry only; skip the soffice/pdftoppm base renders",
    )
    args = parser.parse_args()
    result = extract(Path(args.template), Path(args.outdir), not args.no_base)
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
