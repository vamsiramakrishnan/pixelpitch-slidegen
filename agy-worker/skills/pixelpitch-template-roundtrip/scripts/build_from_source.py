#!/usr/bin/env python3
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     https://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Build a deck INSIDE the customer's actual source template.

The source PPTX is the artifact: selected specimen slides keep their real
masters, artwork, photography and placeholder relationships. This script
only (a) rewrites the copy in their placeholders with shrink-to-fit, and
(b) composes NATIVE visual aids (bar chart / timeline / stat / proportion
bar) on top, in the derived brand palette. Nothing is rasterised; every
added element is an editable PowerPoint shape.

Run with the renderer venv python (python-pptx):
  renderer/.venv/bin/python build_from_source.py \
      --template template.pptx --plan plan.json --out deck.pptx --pngs preview

plan.json:
{
  "palette": {"brand": "#1971ED", "base_dark": "#0E0D26", ...},
  "slides": [
    {
      "specimen": 1,                  // 1-based index into template slides
      "title": "assertion sentence",
      "subtitle": "one line",          // optional
      "body": ["bullet", ...],         // optional, one per body slot
      "visual_aid": {                  // optional; placed below the text
        "type": "bar" | "timeline" | "stat" | "proportion",
        "insight": "annotated takeaway line",
        "unit": "",                    // bar: value suffix
        "series": [{"label": "1924", "value": 1}, ...],   // bar/timeline
        "highlight": "2026",           // bar/timeline: the point
        "value": 96, "denominator": 100,                    // proportion
        "big": "200,000", "caption": "team members"         // stat
      }
    }
  ]
}
"""

from __future__ import annotations

import argparse
import copy
import json
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Emu, Inches, Pt


def _hex(color: str) -> RGBColor:
    return RGBColor.from_string(color.lstrip("#").upper())


def _kind(placeholder) -> str:
    raw = str(placeholder.placeholder_format.type).split()[0].upper()
    return {"TITLE": "TITLE", "SUBTITLE": "SUBTITLE", "BODY": "BODY"}.get(raw, raw)


def _set_text(shape, value: str, *, shrink: bool = True) -> None:
    shape.text = value
    frame = shape.text_frame
    frame.word_wrap = True
    if shrink:
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def _fill_specimen(slide, plan: dict, slide_number: int) -> dict:
    """Rewrite copy in a real specimen slide; keep every visual object."""
    placeholders = list(slide.placeholders)
    titles = sorted(
        [p for p in placeholders if _kind(p) == "TITLE"], key=lambda s: (s.top, s.left)
    )
    subtitles = sorted(
        [p for p in placeholders if _kind(p) == "SUBTITLE" and p.top < Emu(int(Inches(12.3)))],
        key=lambda s: (s.top, s.left),
    )
    bodies = sorted(
        [p for p in placeholders if _kind(p) == "BODY"], key=lambda s: (s.left, s.top)
    )

    filled = []
    consumed = set()
    main_title = None
    if titles:
        if len(titles) > 1:
            counter = next(
                (
                    s
                    for s in titles
                    if (s.has_text_frame and s.text.strip().isdigit())
                    or s.height < Emu(int(Inches(1.0)))
                ),
                titles[0],
            )
            _set_text(counter, f"{slide_number:02d}")
            consumed.add(id(counter))
            filled.append("SECTION_NUMBER")
            main_title = max(
                [s for s in titles if s is not counter], key=lambda s: s.width * s.height
            )
        else:
            main_title = titles[0]
        _set_text(main_title, plan["title"])
        consumed.add(id(main_title))
        filled.append("TITLE")
    if plan.get("subtitle") and subtitles:
        _set_text(subtitles[0], plan["subtitle"])
        consumed.add(id(subtitles[0]))
        filled.append("SUBTITLE")
    queue = list(plan.get("body") or [])
    for slot in bodies:
        if not queue:
            break
        _set_text(slot, queue.pop(0))
        consumed.add(id(slot))
        filled.append("BODY")
    for shape in placeholders:
        if _kind(shape) in {"TITLE", "SUBTITLE", "BODY"} and id(shape) not in consumed:
            _set_text(shape, "")
    return {"filled": filled, "omitted_body": len(queue)}


def _aid_band(slide) -> tuple[int, int, int, int]:
    """Free band below the text: from the lowest text bottom to the footer.

    Uses the union of placeholder boxes on the slide as the text zone and
    leaves a margin above the slide bottom (footer strip).
    """
    bottoms = [p.top + p.height for p in slide.placeholders if p.top is not None]
    top = max(bottoms, default=int(Inches(6.0))) + int(Inches(0.35))
    return top, slide.slide_width, 0, slide.slide_height - int(Inches(1.1))


def _add_aid(slide, aid: dict, palette: dict, band: tuple) -> None:
    top, width, left, bottom = band
    brand = _hex(palette.get("brand", "#1971ED"))
    dark = _hex(palette.get("base_dark", "#0E0D26"))
    rule = _hex(palette.get("rule", "#DEDEDF"))
    fresh = _hex(palette.get("warm_green", "#00B35F"))

    def textbox(x, y, w, h, text, size, color, bold=False, align=1):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        run = frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Arial"
        from pptx.enum.text import PP_ALIGN

        frame.paragraphs[0].alignment = {0: PP_ALIGN.LEFT, 1: PP_ALIGN.CENTER, 2: PP_ALIGN.RIGHT}[align]
        return box

    def bar(x, y, w, h, color):
        from pptx.enum.shapes import MSO_SHAPE

        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        return shape

    height = bottom - top
    top_in, height_in = top / 914400, height / 914400
    kind = aid.get("type")

    if kind == "bar":
        series = aid.get("series") or []
        if not series:
            return
        values = [float(s.get("value", 0)) for s in series]
        vmax = max(values) or 1.0
        chart_w, chart_h = min(13.5, width / 914400 - 2.6), min(4.6, height_in - 1.0)
        base_x = 1.3
        base_y = top_in + chart_h
        n = len(series)
        slot = chart_w / n
        for i, item in enumerate(series):
            h = chart_h * (values[i] / vmax)
            x = base_x + i * slot + slot * 0.18
            w = slot * 0.64
            is_point = str(item.get("label")) == str(aid.get("highlight"))
            bar(x, base_y - h, w, h, brand if is_point else rule)
            textbox(
                x - slot * 0.18, base_y - h - 0.42, w + slot * 0.36, 0.4,
                f"{item.get('value'):,.0f}{aid.get('unit', '')}", 13,
                dark, bold=is_point,
            )
            textbox(x - slot * 0.18, base_y + 0.06, w + slot * 0.36, 0.35,
                    str(item.get("label", "")), 12, dark, bold=is_point)
        if aid.get("insight"):
            textbox(1.3, top_in - 0.05, chart_w, 0.4, aid["insight"], 14, brand, bold=True)

    elif kind == "timeline":
        series = aid.get("series") or []
        if not series:
            return
        from pptx.enum.shapes import MSO_SHAPE

        y = top_in + min(1.7, height_in / 2)
        bar(1.3, y, min(21.0, width / 914400 - 2.6), 0.045, rule)
        n = len(series)
        span = min(21.0, width / 914400 - 2.6)
        step = span / (n - 1) if n > 1 else 0
        for i, item in enumerate(series):
            cx = 1.3 + i * step
            is_point = str(item.get("label")) == str(aid.get("highlight"))
            node = slide.shapes.add_shape(
                MSO_SHAPE.OVAL, Inches(cx - 0.11), Inches(y - 0.085), Inches(0.22), Inches(0.22)
            )
            node.fill.solid()
            node.fill.fore_color.rgb = brand if is_point else dark
            node.line.fill.background()
            above = i % 2 == 0
            label = f"{item.get('label')}"
            if item.get("value"):
                label += f"\n{item.get('value')}{aid.get('unit', '')}"
            textbox(cx - 1.05, y - 0.75 if above else y + 0.22, 2.1, 0.7, label, 12,
                    dark, bold=is_point)
        if aid.get("insight"):
            textbox(1.3, top_in - 0.05, span, 0.4, aid["insight"], 14, brand, bold=True)

    elif kind == "stat":
        big = str(aid.get("big", ""))
        textbox(1.3, top_in + 0.15, 12.0, 1.6, big, 88, brand, bold=True)
        if aid.get("caption"):
            textbox(1.3, top_in + 1.75, 12.0, 0.5, aid["caption"], 18, dark)
        if aid.get("insight"):
            textbox(1.3, top_in + 2.3, 16.0, 0.45, aid["insight"], 15, fresh, bold=True)

    elif kind == "proportion":
        value = float(aid.get("value", 0))
        total = float(aid.get("denominator", 100)) or 100.0
        share = max(0.0, min(1.0, value / total))
        w_total = min(19.0, width / 914400 - 3.0)
        bar(1.3, top_in + 0.35, w_total * share, 1.05, fresh)
        bar(1.3 + w_total * share, top_in + 0.35, w_total * (1 - share), 1.05, rule)
        textbox(1.3, top_in + 1.5, 10.0, 0.45,
                f"{aid.get('value'):,.0f}% {aid.get('caption', '')}", 16, dark, bold=True)
        if aid.get("insight"):
            textbox(1.3, top_in - 0.05, w_total, 0.4, aid["insight"], 14, brand, bold=True)


def build(template: Path, plan: dict, out: Path) -> dict:
    prs = Presentation(str(template))
    slides = list(prs.slides)
    keep_ids = []
    report = []
    for number, spec in enumerate(plan.get("slides", []), start=1):
        index = int(spec.get("specimen", 1)) - 1
        if not (0 <= index < len(slides)):
            raise SystemExit(f"plan references specimen {index + 1}; template has {len(slides)}")
        slide = slides[index]
        fill = _fill_specimen(slide, spec, number)
        aid_added = bool(spec.get("visual_aid"))
        if aid_added:
            _add_aid(slide, spec["visual_aid"], plan.get("palette") or {}, _aid_band(slide))
        keep_ids.append((index, slide))
        report.append({"specimen": index + 1, "aid": (spec.get("visual_aid") or {}).get("type"), **fill})

    # Keep only chosen specimens, in narrative order.
    id_map = {id(slides[i]): i for i in range(len(slides))}
    chosen = [slide for _, slide in keep_ids]
    xml_slides = prs.slides._sldIdLst
    slide_ids = list(xml_slides)
    chosen_ids = [slide_ids[id_map[id(s)]] for s in chosen]
    for sldid in slide_ids:
        if sldid not in chosen_ids:
            prs.part.drop_rel(sldid.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
            xml_slides.remove(sldid)
    for sldid in chosen_ids:
        xml_slides.remove(sldid)
        xml_slides.append(sldid)

    prs.save(str(out))
    return {"slides": report, "slide_count": len(chosen)}


def rasterize(pptx: Path, outdir: Path) -> list[Path]:
    outdir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(pptx)],
        check=True, capture_output=True, timeout=420,
    )
    pdf = next(iter(outdir.glob("*.pdf")), None)
    if pdf is None:
        raise SystemExit("pptx did not convert to pdf")
    subprocess.run(
        ["pdftoppm", "-png", "-r", "110", str(pdf), str(outdir / "slide")],
        check=True, capture_output=True, timeout=420,
    )
    return sorted(outdir.glob("slide*.png"))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", required=True)
    parser.add_argument("--plan", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--pngs", help="also rasterize to this directory")
    parser.add_argument("--report", help="write the build report json here")
    args = parser.parse_args()

    plan = json.loads(Path(args.plan).read_text(encoding="utf-8"))
    result = build(Path(args.template), plan, Path(args.out))
    if args.pngs:
        result["pngs"] = [str(p) for p in rasterize(Path(args.out), Path(args.pngs))]
    payload = json.dumps(result, indent=2)
    if args.report:
        Path(args.report).write_text(payload, encoding="utf-8")
    print(payload)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
